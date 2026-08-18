"""
检测报告 → 功效摘要 Word

- 有文字层 PDF：直接抽字（优先结论页）
- MiniMax-M3：摘抄成结构化摘要（对照样例 docx 风格）
- python-docx 导出
"""
from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

import pymupdf
from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Pt, Cm

from app import minimax


def extract_pdf_text(
    pdf_path: str | Path,
    *,
    start_page: int | None = None,
    end_page: int | None = None,
    max_chars: int = 28000,
) -> dict[str, Any]:
    """
    抽取 PDF 文本。若未指定页码，自动定位「检测结论」附近。
    页码为 1-based inclusive。
    """
    doc = pymupdf.open(pdf_path)
    n = doc.page_count
    # 优先「章节式」结论标题；避免目录/前言里的零散命中
    strong_hits: list[int] = []
    weak_hits: list[int] = []
    for i in range(n):
        t = doc[i].get_text("text") or ""
        if re.search(r"(?:^|\n)\s*(?:\d+[\.、]\s*)?检测结论", t):
            strong_hits.append(i + 1)
        elif "检测结论" in t:
            weak_hits.append(i + 1)
        elif "综上所述" in t and i + 1 > n * 0.4:
            # 文末总结
            weak_hits.append(i + 1)

    conclusion_page = None
    if strong_hits:
        # 取最后一次强命中（通常是正文结论，不是目录）
        conclusion_page = strong_hits[-1]
    elif weak_hits:
        conclusion_page = weak_hits[-1]

    if start_page is None:
        if conclusion_page:
            start_page = max(1, conclusion_page)
        else:
            # 默认后 1/3
            start_page = max(1, int(n * 0.65))
    if end_page is None:
        if conclusion_page:
            end_page = min(n, conclusion_page + 6)
        else:
            end_page = n

    start_page = max(1, min(start_page, n))
    end_page = max(start_page, min(end_page, n))

    parts: list[str] = []
    pages_used: list[int] = []
    for i in range(start_page - 1, end_page):
        t = (doc[i].get_text("text") or "").strip()
        if not t:
            continue
        pages_used.append(i + 1)
        parts.append(f"===== 第 {i+1}/{n} 页 =====\n{t}")
    doc.close()

    full = "\n\n".join(parts)
    if len(full) > max_chars:
        full = full[:max_chars] + "\n…(截断)"

    # 报告编号
    report_no = ""
    m = re.search(r"报告编号[：:]\s*([A-Za-z0-9\-]+)", full)
    if m:
        report_no = m.group(1)

    product = ""
    m2 = re.search(r"试验产品[「\"']?([^」\"'\n]{4,40})", full)
    if m2:
        product = m2.group(1).strip()

    return {
        "page_count": n,
        "start_page": start_page,
        "end_page": end_page,
        "pages_used": pages_used,
        "conclusion_page": conclusion_page,
        "report_no": report_no,
        "product_guess": product,
        "text": full,
        "char_count": len(full),
    }


SUMMARY_SYSTEM = """你是化妆品功效检测报告摘抄员。根据检测报告结论页文本，生成「功效摘要」文稿。

要求：
1. 摘抄，不要编造报告中没有的数字或功效。
2. 风格对照业务样例：分点 1、2、… 写受试者概况 + 结论句 + 关键数据要点。
3. 若报告含「单品」与「水乳组合」两套试验，分别写两点；只有一套则写一点。
4. 必须只输出 JSON：
{
  "title": "产品名 + 功效摘要",
  "product_name": "...",
  "report_no": "...",
  "paragraphs": [
    "1、……完整段落……",
    "2、……"
  ],
  "claims": ["保湿","修护",...],
  "key_metrics": [
    {"name":"角质层水分含量 28天","value":"上升率36.84%","note":"p<0.001"}
  ],
  "safety": "温和无刺激/敏感肌适用等一句话"
}
"""


def summarize_with_minimax(extract: dict[str, Any]) -> dict[str, Any]:
    user = {
        "report_no": extract.get("report_no"),
        "product_guess": extract.get("product_guess"),
        "pages": f"{extract.get('start_page')}-{extract.get('end_page')}",
        "source_text": extract.get("text") or "",
    }
    messages = [
        {"role": "system", "content": SUMMARY_SYSTEM},
        {
            "role": "user",
            "content": "请根据下列检测报告结论摘抄功效摘要，只输出 JSON：\n"
            + json.dumps(user, ensure_ascii=False),
        },
    ]
    text, meta = minimax.chat(messages, temperature=0.15, max_tokens=4096)
    parsed = minimax._extract_json(text)  # noqa: SLF001
    if not isinstance(parsed, dict):
        # 兜底：把原文压缩成两段
        parsed = {
            "title": (extract.get("product_guess") or "检测报告") + "功效摘要",
            "product_name": extract.get("product_guess") or "",
            "report_no": extract.get("report_no") or "",
            "paragraphs": [
                "（模型未返回结构化 JSON，以下为结论页原文摘录，请人工整理）",
                (extract.get("text") or "")[:2500],
            ],
            "claims": [],
            "key_metrics": [],
            "safety": "",
        }
    parsed["_meta"] = meta
    parsed["_raw"] = text[:3000]
    return parsed


def build_docx(summary: dict[str, Any], out_path: str | Path) -> Path:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(12)
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(8)

    title = summary.get("title") or "功效摘要"
    h = doc.add_heading(title, level=1)
    for run in h.runs:
        run.font.size = Pt(16)

    meta_bits = []
    if summary.get("product_name"):
        meta_bits.append(f"产品：{summary['product_name']}")
    if summary.get("report_no"):
        meta_bits.append(f"报告编号：{summary['report_no']}")
    if meta_bits:
        p = doc.add_paragraph(" · ".join(meta_bits))
        for run in p.runs:
            run.font.size = Pt(10)
            run.font.color.rgb = None

    for para in summary.get("paragraphs") or []:
        t = str(para).strip()
        if t:
            doc.add_paragraph(t)

    claims = summary.get("claims") or []
    if claims:
        doc.add_heading("功效宣称要点", level=2)
        doc.add_paragraph("、".join(str(c) for c in claims))

    metrics = summary.get("key_metrics") or []
    if metrics:
        doc.add_heading("关键数据", level=2)
        for m in metrics:
            if isinstance(m, dict):
                line = f"{m.get('name','')}：{m.get('value','')}"
                if m.get("note"):
                    line += f"（{m['note']}）"
            else:
                line = str(m)
            doc.add_paragraph(line, style="List Bullet")

    if summary.get("safety"):
        doc.add_heading("安全性", level=2)
        doc.add_paragraph(str(summary["safety"]))

    foot = doc.add_paragraph(
        "本摘要由备案审核工作台自检测报告结论页摘抄生成，AI 辅助整理，请人工核对原文后使用。"
    )
    for run in foot.runs:
        run.font.size = Pt(9)

    doc.save(str(out_path))
    return out_path


def run_report_summary_job(
    tid: str,
    pdf_path: Path,
    out_dir: Path,
    *,
    title: str = "检测报告功效摘要",
    start_page: int | None = None,
    end_page: int | None = None,
    use_ai: bool = True,
) -> dict[str, Any]:
    extract = extract_pdf_text(pdf_path, start_page=start_page, end_page=end_page)
    summary: dict[str, Any]
    engine = "pdf-text"
    if use_ai:
        try:
            summary = summarize_with_minimax(extract)
            engine = f"pdf-text + {(summary.get('_meta') or {}).get('model') or 'MiniMax'}"
        except Exception as e:
            summary = {
                "title": title,
                "product_name": extract.get("product_guess") or "",
                "report_no": extract.get("report_no") or "",
                "paragraphs": [
                    f"（AI 摘要失败：{e}。以下为结论页原文摘录）",
                    (extract.get("text") or "")[:3000],
                ],
                "claims": [],
                "key_metrics": [],
                "safety": "",
                "error": str(e),
            }
            engine = "pdf-text (AI failed)"
    else:
        summary = {
            "title": title,
            "product_name": extract.get("product_guess") or "",
            "report_no": extract.get("report_no") or "",
            "paragraphs": [(extract.get("text") or "")[:3000]],
            "claims": [],
            "key_metrics": [],
            "safety": "",
        }

    if not summary.get("title"):
        summary["title"] = title
    docx_name = "efficacy_summary.docx"
    docx_path = out_dir / docx_name
    build_docx(summary, docx_path)

    # 伪 hits 便于统一任务 UI / 审计
    hits = [
        {
            "id": "summary",
            "field": "功效摘要",
            "excel_value": "\n\n".join(summary.get("paragraphs") or [])[:2000],
            "status": "一致" if not summary.get("error") else "疑点",
            "evidence": f"摘抄页 {extract.get('start_page')}-{extract.get('end_page')} · {engine}",
            "score": 100.0 if not summary.get("error") else 50.0,
            "decision": "pending",
            "bboxes": [],
            "page": extract.get("conclusion_page") or extract.get("start_page") or 1,
            "category": "summary",
        }
    ]
    for i, c in enumerate(summary.get("claims") or []):
        hits.append(
            {
                "id": f"claim_{i}",
                "field": f"功效 · {c}",
                "excel_value": str(c),
                "status": "一致",
                "evidence": "模型摘抄宣称点",
                "score": 100.0,
                "decision": "pending",
                "bboxes": [],
                "page": 1,
                "category": "claim",
            }
        )

    return {
        "type": "report_summary",
        "engine": engine,
        "extract": {
            "page_count": extract["page_count"],
            "start_page": extract["start_page"],
            "end_page": extract["end_page"],
            "pages_used": extract["pages_used"],
            "conclusion_page": extract["conclusion_page"],
            "report_no": extract["report_no"],
            "char_count": extract["char_count"],
        },
        "summary": summary,
        "docx_name": docx_name,
        "docx_url": f"/api/tasks/{tid}/docx",
        "hits": hits,
        "summary_counts": {
            "一致": sum(1 for h in hits if h["status"] == "一致"),
            "疑点": sum(1 for h in hits if h["status"] == "疑点"),
            "缺失": 0,
            "跳过": 0,
        },
        "ocr_text": (extract.get("text") or "")[:8000],
        "pages": [],  # 报告任务不强制渲染图
    }
