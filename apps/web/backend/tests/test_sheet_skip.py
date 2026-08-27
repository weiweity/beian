from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
import pytest

from app.fields import parse_excel_fields
from app.layout_zones import skip_sheet_field


def test_skip_sheet_field_names():
    assert skip_sheet_field("工艺说明")
    assert skip_sheet_field("颜色要求")
    assert skip_sheet_field("版本号")
    assert skip_sheet_field("  版本号 ")
    assert skip_sheet_field("更新内容：调整净含量")
    assert not skip_sheet_field("中文品名")
    assert not skip_sheet_field("生产信息及其他")
    assert not skip_sheet_field("备案版本号")
    assert not skip_sheet_field("执行标准版本号")


def test_parse_excel_skips_process_rows(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append([1, "中文品名", "达肤妍祛痘细肤面膜", None, ""])
    ws.append([2, "工艺说明", "烫金", None, ""])
    ws.append([3, "颜色要求", "专色", None, ""])
    ws.append([4, "版本号", "26H06A", None, ""])
    ws.append([5, "更新内容", "净含量改为7片", None, ""])
    ws.append([6, "净含量", "7片", None, ""])
    path = tmp_path / "confirm.xlsx"
    wb.save(path)
    fields = parse_excel_fields(str(path))
    names = [f["field"] for f in fields]
    assert names == ["中文品名", "净含量"]


def test_parse_excel_ignores_hidden_remark_column(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append([1, "中文品名", "达肤妍海葡萄油萃微珠保湿喷雾", None, "⠓⠪⠏⠥⠞⠖⠺⠌⠥⠃⠖⠱⠁⠏⠴⠥"])
    ws.column_dimensions["E"].hidden = True
    path = tmp_path / "hidden-remark.xlsx"
    wb.save(path)

    fields = parse_excel_fields(str(path))

    assert fields == [
        {
            "field": "中文品名",
            "excel_value": "达肤妍海葡萄油萃微珠保湿喷雾",
            "remark": "",
            "field_group": "中文品名",
        }
    ]


def test_parse_excel_ignores_hidden_rows(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append([1, "中文品名", "达肤妍海葡萄油萃微珠保湿喷雾"])
    ws.append([2, "内部辅助字段", "不要进入机审"])
    ws.row_dimensions[2].hidden = True
    path = tmp_path / "hidden-row.xlsx"
    wb.save(path)

    fields = parse_excel_fields(str(path))

    assert [field["field"] for field in fields] == ["中文品名"]


def test_parse_excel_rejects_when_all_business_rows_are_hidden(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append([1, "中文品名", "达肤妍海葡萄油萃微珠保湿喷雾"])
    ws.row_dimensions[1].hidden = True
    path = tmp_path / "all-business-rows-hidden.xlsx"
    wb.save(path)

    with pytest.raises(ValueError, match="没有可核对字段"):
        parse_excel_fields(str(path))


def test_parse_excel_stops_at_confirmation_footer(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append([1, "生产信息及其他", "生产商：伸美"])
    ws.append([2, "职责", "包装设计部"])
    ws.append([3, "设计部联系人", "不进入机审"])
    ws.append([4, "签字", "审批签字"])
    path = tmp_path / "footer-boundary.xlsx"
    wb.save(path)

    fields = parse_excel_fields(str(path))

    assert [field["field"] for field in fields] == ["生产信息及其他"]


def test_parse_excel_rejects_footer_only_sheet(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append([1, "职责", "包装设计部"])
    ws.append([2, "签字", "审批签字"])
    path = tmp_path / "footer-only.xlsx"
    wb.save(path)

    with pytest.raises(ValueError, match="没有可核对字段"):
        parse_excel_fields(str(path))


def test_parse_excel_ignores_sheet_when_required_contract_column_is_hidden(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append([1, "中文品名", "不应进入机审"])
    ws.column_dimensions["C"].hidden = True
    path = tmp_path / "hidden-required-column.xlsx"
    wb.save(path)

    with pytest.raises(ValueError, match="项目列或内容列被隐藏"):
        parse_excel_fields(str(path))


def test_parse_excel_uses_visible_d_remark_and_promotes_it_when_value_is_empty(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append([1, "中文品名", "", "达肤妍海葡萄油萃微珠保湿喷雾", ""])
    path = tmp_path / "visible-d-remark.xlsx"
    wb.save(path)

    assert parse_excel_fields(str(path)) == [
        {
            "field": "中文品名",
            "excel_value": "达肤妍海葡萄油萃微珠保湿喷雾",
            "remark": "",
            "field_group": "中文品名",
        }
    ]


def test_parse_excel_still_splits_visible_ingredient_steps(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append([1, "成分表", "步骤01：水、甘油\n步骤02：海葡萄提取物"])
    path = tmp_path / "ingredient-steps.xlsx"
    wb.save(path)

    fields = parse_excel_fields(str(path))

    assert [field["step_id"] for field in fields] == ["step_1", "step_2"]
    assert all(field["field_group"] == "成分表" for field in fields)
